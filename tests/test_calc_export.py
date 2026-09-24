"""calc_export 导出单元测试（阶段5 C 方案批次3）— 内存库 + openpyxl 回读校验。

覆盖：
- `analyze_formula`：DATA/PARAM 判定、**`+#REF!` 伪节点必须填值（N2）**、
  跨表引用集合收集、解析失败保守处理。
- `should_keep_formula`：跨表引用**只在被引用表同在导出集合内**时才保留。
- `_safe_title`：Excel 标题硬约束（≤31 字符 / 禁用字符替换 / 空名兜底）。
- **多表闭包导出**：目标表 + 它递归引用到的表同簿写出，跨表公式保留且引号正确，
  导出的公式能被本引擎再解析（往返自洽）。
- `#REF!` 格子写**文本**而非公式（写成公式 Excel 打开即报错）。
- 仅值版：一律写计算值。
- `_apply_drift`：老数据标题漂移时同步改写公式里的表名。

运行：python tests/test_calc_export.py
"""
import sqlite3
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.db import SCHEMA  # noqa: E402
from app.engine import calc_sheet as cs  # noqa: E402
from app.engine.calc_ast_render import render  # noqa: E402
from app.engine.calc_formula import Parser, ParseError  # noqa: E402
from app.exporter.calc_export import (  # noqa: E402
    _apply_drift, _safe_title, analyze_formula, export_sheet, should_keep_formula,
)

OK, FAILS = 0, []


def check(label, cond, detail=""):
    global OK
    if cond:
        OK += 1
    else:
        FAILS.append(f"{label} {detail}".strip())


def roundtrip(body: str) -> str:
    """公式体 → AST → 渲染，用于验证导出的公式能被本引擎原样再解析。"""
    return render(Parser(body).parse())


def main() -> int:
    # ===== 纯逻辑：analyze_formula =====
    h, e, s = analyze_formula("=A1*2")
    check("纯公式 无宿主函数", h is False)
    check("纯公式 无 #REF!", e is False)
    check("纯公式 无跨表", s == set(), f"got={s}")

    h, e, _ = analyze_formula('=DATA("周立生","业务收入",2025,3)')
    check("DATA 判宿主函数", h is True)

    h, _, _ = analyze_formula('=PARAM("k")')
    check("PARAM 判宿主函数", h is True)

    h, _, _ = analyze_formula('=A1+DATA("x","y",2025)')
    check("混合公式判宿主函数", h is True)

    h, _, _ = analyze_formula('=SUM(A1:A10)*PARAM("k")')
    check("函数内混合判宿主函数", h is True)

    # --- N2：#REF! 伪节点（G2 删行列产物）写成公式会让 Excel 打不开 ---
    _, e, _ = analyze_formula("=#REF!")
    check("N2 裸 #REF! 判失效", e is True)
    _, e, _ = analyze_formula("=A1+#REF!")
    check("N2 混合 #REF! 判失效", e is True)
    _, e, _ = analyze_formula("=A1+B2")
    check("正常公式不误判失效", e is False)

    # --- 跨表引用集合（含引号名）---
    _, _, s = analyze_formula("='2026-01'!A1+1")
    check("引号表名取出真名", s == {"2026-01"}, f"got={s}")
    _, _, s = analyze_formula("=SUM('2025分成'!A1:B2)")
    check("引号表名在函数内", s == {"2025分成"}, f"got={s}")
    _, _, s = analyze_formula("=甲表!A1+乙表!B2")
    check("多个普通表名", s == {"甲表", "乙表"}, f"got={s}")
    _, _, s = analyze_formula("='a''b'!A1")
    check("转义引号还原", s == {"a'b"}, f"got={s}")

    h, e, s = analyze_formula("=1+")
    check("解析失败保守填值", h is True and e is True and s == set(), f"got={h},{e},{s}")

    # ===== 纯逻辑：should_keep_formula =====
    check("纯公式保留", should_keep_formula("=A1*2") is True)
    check("SUM 保留", should_keep_formula("=SUM(A1:A10)") is True)
    check("DATA 填值", should_keep_formula('=DATA("x","y",2025)') is False)
    check("PARAM 填值", should_keep_formula('=PARAM("k")') is False)
    check("N2 #REF! 不写公式", should_keep_formula("=#REF!") is False)
    check("非公式不保留", should_keep_formula("100") is False)
    check("空串不保留", should_keep_formula("") is False)
    # 跨表：被引用表在集合内 → 保留；不在 → 填值（否则 Excel 里就是 #REF!）
    check("跨表 被引用表在集合内保留",
          should_keep_formula("='2026-01'!A1", {"2026-01"}) is True)
    check("跨表 被引用表不在集合内填值",
          should_keep_formula("='2026-01'!A1", {"主表"}) is False)
    check("跨表 known_sheets=None 放行",
          should_keep_formula("='2026-01'!A1") is True)

    # ===== 纯逻辑：_safe_title（Excel 标题硬约束）=====
    check("标题原样", _safe_title("主表") == "主表")
    check("标题含 '-' 原样", _safe_title("2026-01") == "2026-01")
    check("标题替换禁用字符", _safe_title("a/b") == "a_b", f"got={_safe_title('a/b')}")
    check("标题截断到 31", len(_safe_title("x" * 50)) == 31)
    check("标题空名兜底", _safe_title("") == "Sheet1")

    # ===== 纯逻辑：_apply_drift（老数据标题漂移兜底）=====
    drift = {"2025年度分成计算表超长名称abcdef": "2025年度分成计算表超长名称abc"}
    check("漂移改写引号名",
          _apply_drift("='2025年度分成计算表超长名称abcdef'!A1", drift)
          == "='2025年度分成计算表超长名称abc'!A1",
          f"got={_apply_drift(chr(39) + '2025年度分成计算表超长名称abcdef' + chr(39) + '!A1', drift)}")
    check("漂移不影响普通名", _apply_drift("=甲表!A1", drift) == "=甲表!A1")
    check("无漂移原样", _apply_drift("='2026-01'!A1", {}) == "='2026-01'!A1")

    # ===== DB + openpyxl：多表闭包导出 =====
    tmp = Path(tempfile.mkdtemp(prefix="calc_exp_"))
    db = tmp / "exp.db"
    conn = sqlite3.connect(db)
    conn.row_factory = sqlite3.Row
    conn.executescript(SCHEMA)
    conn.commit()

    # 三级链：主表 → 2026-01（需引号）→ 三级表（普通名，不需引号）
    sid3 = cs.create_sheet("三级表", conn=conn)
    cs.save_content(sid3, cs.default_content(3, 2) | {
        "cells": {"0,0": {"raw": "42", "kind": "number"}}}, conn=conn)

    sid2 = cs.create_sheet("2026-01", conn=conn)
    cs.save_content(sid2, cs.default_content(3, 2) | {
        "cells": {
            "0,0": {"raw": "5", "kind": "number"},
            "1,0": {"raw": "=三级表!A1*2", "kind": "formula"},   # 普通名无需引号
        }}, conn=conn)

    sid1 = cs.create_sheet("主表", conn=conn)
    cs.save_content(sid1, cs.default_content(6, 2) | {
        "cells": {
            "0,0": {"raw": "100", "kind": "number"},
            "1,0": {"raw": "=A1*2", "kind": "formula"},
            "2,0": {"raw": "='2026-01'!A1+1", "kind": "formula"},   # 数字开头 → 引号
            "3,0": {"raw": "=#REF!", "kind": "formula"},            # N2 → 必须填值
        }}, conn=conn)

    # 悬空引用：表名在库里不存在 → 不在闭包内 → 填值而非写断链公式
    sid4 = cs.create_sheet("孤表", conn=conn)
    cs.save_content(sid4, cs.default_content(3, 2) | {
        "cells": {"0,0": {"raw": "=幽灵表!A1", "kind": "formula"}}}, conn=conn)

    f1 = export_sheet(sid1, str(tmp / "带公式.xlsx"), with_formula=True, conn=conn)
    from openpyxl import load_workbook  # noqa: E402
    wb = load_workbook(f1)

    check("闭包导出 3 张表", set(wb.sheetnames) == {"主表", "2026-01", "三级表"},
          f"got={wb.sheetnames}")
    check("目标表排第一", wb.sheetnames[0] == "主表", f"got={wb.sheetnames[0]}")

    ws = wb["主表"]
    check("主表 A1 数字", ws["A1"].value == 100, f"got={ws['A1'].value}")
    check("主表 A2 保留公式", ws["A2"].value == "=A1*2", f"got={ws['A2'].value}")
    check("跨表公式保留且带引号", ws["A3"].value == "='2026-01'!A1+1",
          f"got={ws['A3'].value}")
    v = ws["A4"].value
    check("N2 #REF! 写文本不是公式",
          v == "#REF!" and not (isinstance(v, str) and v.startswith("=")),
          f"got={v!r}")

    # 导出的公式必须能被本引擎原样再解析（往返自洽）
    try:
        check("导出公式可再解析", roundtrip("'2026-01'!A1+1") == "'2026-01'!A1+1",
              f"got={roundtrip(chr(39) + '2026-01' + chr(39) + '!A1+1')}")
    except ParseError as ex:  # pragma: no cover
        check("导出公式可再解析", False, f"ParseError={ex}")

    ws2 = wb["2026-01"]
    check("被引用表 A1 值", ws2["A1"].value == 5, f"got={ws2['A1'].value}")
    check("被引用表内公式保留", ws2["A2"].value == "=三级表!A1*2", f"got={ws2['A2'].value}")
    check("第三级表值", wb["三级表"]["A1"].value == 42,
          f"got={wb['三级表']['A1'].value}")

    # 悬空引用 → 填值（不能写成 Excel 断链公式）
    f2 = export_sheet(sid4, str(tmp / "孤表.xlsx"), with_formula=True, conn=conn)
    v4 = load_workbook(f2)["孤表"]["A1"].value
    check("悬空跨表引用填值",
          not (isinstance(v4, str) and v4.startswith("=")), f"got={v4!r}")

    # ===== 仅值版 =====
    f3 = export_sheet(sid1, str(tmp / "仅值.xlsx"), with_formula=False, conn=conn)
    wsv = load_workbook(f3)["主表"]
    check("仅值 A1", wsv["A1"].value == 100, f"got={wsv['A1'].value}")
    check("仅值 A2=A1*2", wsv["A2"].value == 200, f"got={wsv['A2'].value}")
    check("仅值 A3=跨表值", wsv["A3"].value == 6, f"got={wsv['A3'].value}")
    check("仅值 A4 不是公式",
          not (isinstance(wsv["A4"].value, str) and wsv["A4"].value.startswith("=")),
          f"got={wsv['A4'].value!r}")

    conn.close()
    return 1 if FAILS else 0


if __name__ == "__main__":
    rc = main()
    print(f"PASS {OK}" if not FAILS else "FAILED:\n" + "\n".join(FAILS))
    sys.exit(rc)
