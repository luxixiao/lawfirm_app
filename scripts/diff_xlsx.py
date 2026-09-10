#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""diff_xlsx.py — 逐格跨 sheet 比较两个 xlsx 文件（golden 回归安全网）

用途
----
C#/WPF 重写期，把 C# 导出的 xlsx 与 Python 应用导出的 golden 基准逐格 diff，
保证不回归。脚本**纯 openpyxl、无 UI 依赖、可无头运行**，不依赖 PySide6。

比对维度（覆盖计划 T1 验证项 ① 要求的列名/合并/数字格式/页面设置逐格 diff）：
  1. sheet 名 / 顺序
  2. 每 sheet 列名序列（自动探测表头行，比较其单元格值序列）
  3. 合并单元格区域
  4. 合计 / 汇总 / 结余 行（按行扫描 合计/汇总/结余 关键词，逐格比较）
  5. 数字格式（每格 number_format）
  6. 页面设置（orientation / fitToWidth / fitToHeight / 四边 margins）
  7. 逐格比对（覆盖所有数据格、表头、合计行、表尾，含值 + 数字格式）

数值比较采用容差（默认 1e-6），以容忍 C# decimal 与 Python float 的末位差异；
其余类型按 `==` 比较。任一维度不一致即判定 FAIL，进程退出码非 0。

用法
----
    python scripts/diff_xlsx.py <baseline.xlsx> <candidate.xlsx> [--json out.json] [--verbose] [--tol 1e-6]

退出码
------
    0  两侧完全一致
    1  存在任意不一致
    2  参数错误 / 文件无法读取
"""
from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

# 合计/汇总/结余 行关键词（用于自动定位汇总行）
_SUMMARY_KEYWORDS = ("合计", "汇总", "结余")


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------
@dataclass
class Diff:
    """单条不一致记录。"""

    category: str               # sheet / cell / number_format / merged / header / summary_row / page_setup
    sheet: str                  # 所在 sheet 标题（跨 sheet 维度为 "<workbook>"）
    location: str               # 坐标 / 范围 / 属性名
    expected: Any               # baseline 侧值
    actual: Any                 # candidate 侧值
    note: str = ""              # 补充说明

    def to_dict(self) -> Dict[str, Any]:
        return {
            "category": self.category,
            "sheet": self.sheet,
            "location": self.location,
            "expected": self.expected,
            "actual": self.actual,
            "note": self.note,
        }


@dataclass
class DiffReport:
    """完整比对报告。"""

    baseline: str = ""
    candidate: str = ""
    tol: float = 1e-6
    diffs: List[Diff] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return len(self.diffs) == 0

    def add(self, *args, **kwargs) -> None:
        self.diffs.append(Diff(*args, **kwargs))

    def to_dict(self) -> Dict[str, Any]:
        cats: Dict[str, int] = {}
        for d in self.diffs:
            cats[d.category] = cats.get(d.category, 0) + 1
        return {
            "baseline": self.baseline,
            "candidate": self.candidate,
            "tolerance": self.tol,
            "passed": self.passed,
            "diff_count": len(self.diffs),
            "by_category": cats,
            "diffs": [d.to_dict() for d in self.diffs],
        }


# ---------------------------------------------------------------------------
# 取值 / 比较辅助
# ---------------------------------------------------------------------------
def _values_equal(a: Any, b: Any, tol: float) -> bool:
    """比较两个单元格值。数值按容差比较，其余按相等比较。"""
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) <= tol
    # datetime / str / None 等按原生相等
    return a == b


def _cell_map(ws: "openpyxl.worksheet.worksheet.Worksheet") -> Dict[Tuple[int, int], Tuple[Any, str]]:
    """采集 sheet 内所有 (行,列) -> (值, 数字格式)。"""
    out: Dict[Tuple[int, int], Tuple[Any, str]] = {}
    max_r = ws.max_row or 0
    max_c = ws.max_column or 0
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            cell = ws.cell(row=r, column=c)
            out[(r, c)] = (cell.value, cell.number_format or "General")
    return out


def _detect_header_row(ws: "openpyxl.worksheet.worksheet.Worksheet") -> Optional[int]:
    """探测表头行：前 5 行中字符串单元格最多的一行（且 >=2 个字符串）。"""
    best, best_count = None, 0
    for r in range(1, min(ws.max_row or 0, 5) + 1):
        count = 0
        for c in range(1, (ws.max_column or 0) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip():
                count += 1
        if count > best_count:
            best, best_count = r, count
    return best if best_count >= 2 else None


def _header_sequence(ws: "openpyxl.worksheet.worksheet.Worksheet") -> List[Any]:
    """表头行的值序列（左→右）。"""
    hrow = _detect_header_row(ws)
    if hrow is None:
        return []
    return [ws.cell(row=hrow, column=c).value for c in range(1, (ws.max_column or 0) + 1)]


def _summary_rows(ws: "openpyxl.worksheet.worksheet.Worksheet") -> Dict[int, List[Any]]:
    """扫描含合计/汇总/结余 关键词的行，返回 {行号: 整行值列表}。"""
    rows: Dict[int, List[Any]] = {}
    for r in range(1, (ws.max_row or 0) + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, (ws.max_column or 0) + 1)]
        labels = [str(v) for v in vals if isinstance(v, str)]
        if any(any(kw in lb for kw in _SUMMARY_KEYWORDS) for lb in labels):
            rows[r] = vals
    return rows


def _merged_list(ws: "openpyxl.worksheet.worksheet.Worksheet") -> List[str]:
    """合并区域排序列表（如 ['A1:O1', 'C3:D3']）。"""
    return sorted(str(r) for r in ws.merged_cells.ranges)


def _fmt(v: Any) -> str:
    """用于文本报告的值格式化。"""
    if v is None:
        return "<empty>"
    if isinstance(v, float):
        return repr(v)
    return str(v)


# ---------------------------------------------------------------------------
# 核心比对
# ---------------------------------------------------------------------------
def compare_workbooks(base_wb: "openpyxl.Workbook", cand_wb: "openpyxl.Workbook",
                       tol: float, report: DiffReport) -> None:
    """逐维度比对两个 workbook，结果写入 report。"""
    base_sheets = [ws.title for ws in base_wb.worksheets]
    cand_sheets = [ws.title for ws in cand_wb.worksheets]

    # ---- 1. sheet 名 / 顺序 ----
    base_set, cand_set = set(base_sheets), set(cand_sheets)
    for t in base_sheets:
        if t not in cand_set:
            report.add("sheet", "<workbook>", t, "present", "absent", "baseline 存在但 candidate 缺失")
    for t in cand_sheets:
        if t not in base_set:
            report.add("sheet", "<workbook>", t, "absent", "present", "candidate 多出 baseline 没有的 sheet")
    for i, (a, b) in enumerate(zip(base_sheets, cand_sheets)):
        if a != b:
            report.add("sheet", "<workbook>", f"#{i}", a, b, "sheet 顺序不一致")

    # ---- 逐 sheet 比对 ----
    common = [t for t in base_sheets if t in cand_set]
    for title in common:
        bws = base_wb[title]
        cws = cand_wb[title]
        _compare_sheet(bws, cws, title, tol, report)


def _compare_sheet(bws, cws, title: str, tol: float, report: DiffReport) -> None:
    """比对单个 sheet 的 逐格 / 数字格式 / 合并 / 表头 / 汇总行 / 页面设置。"""
    # ---- 2. 逐格 + 数字格式 ----
    b_cells = _cell_map(bws)
    c_cells = _cell_map(cws)
    coords = set(b_cells) | set(c_cells)
    for coord in sorted(coords):
        (bv, bf), (cv, cf) = b_cells.get(coord, (None, "General")), c_cells.get(coord, (None, "General"))
        if not _values_equal(bv, cv, tol):
            loc = f"{get_column_letter(coord[1])}{coord[0]}"
            report.add("cell", title, loc, bv, cv, "单元格值不一致")
        # 数字格式仅在有值的一侧存在时比较（跳过纯空单元格的格式噪音）
        if (bv is not None or cv is not None) and bf != cf:
            loc = f"{get_column_letter(coord[1])}{coord[0]}"
            report.add("number_format", title, loc, bf, cf, "数字格式不一致")

    # ---- 3. 合并单元格区域 ----
    bm, cm = _merged_list(bws), _merged_list(cws)
    for r in bm:
        if r not in set(cm):
            report.add("merged", title, r, "present", "absent", "baseline 合并区域在 candidate 缺失")
    for r in cm:
        if r not in set(bm):
            report.add("merged", title, r, "absent", "present", "candidate 多出 baseline 没有的合并区域")

    # ---- 4. 列名序列（表头行）----
    bh, ch = _header_sequence(bws), _header_sequence(cws)
    if bh or ch:
        if bh != ch:
            report.add("header", title, f"row {_detect_header_row(bws)}", bh, ch, "列名序列不一致")

    # ---- 5. 合计/汇总/结余 行 ----
    bs, cs = _summary_rows(bws), _summary_rows(cws)
    for r in sorted(set(bs) | set(cs)):
        if r not in cs:
            report.add("summary_row", title, f"row {r}", bs[r], None, "汇总行在 candidate 缺失")
        elif r not in bs:
            report.add("summary_row", title, f"row {r}", None, cs[r], "candidate 多出汇总行")
        elif bs[r] != cs[r]:
            report.add("summary_row", title, f"row {r}", bs[r], cs[r], "汇总行数值不一致")

    # ---- 6. 页面设置 ----
    _compare_page_setup(bws, cws, title, report, tol)


def _compare_page_setup(bws, cws, title: str, report: DiffReport, tol: float = 1e-6) -> None:
    """比对 orientation / fitToWidth / fitToHeight / 四边 margins。

    归一化说明：
    * fitToWidth / fitToHeight 在 SpreadsheetML 里 schema 默认值就是 1。openpyxl 会显式写
      fitToWidth="1"，而 NPOI 的 CT_PageSetup 带 [DefaultValue(1)]、XmlSerializer 在等于
      默认值时**省略**该属性，读回来是 None。二者对 Excel 语义完全等价，因此把 None 视为 1。
    * 页边距按 tol 做浮点容差比较，避免 0.3 / 0.30000001192092896 这类二进制表示噪声。
    """
    bp, cp = bws.page_setup, cws.page_setup

    def fit(v):
        # schema 默认值归一：None（未落盘）== 1
        return 1 if v is None else v

    attrs = [
        ("orientation", lambda p: p.orientation),
        ("fitToWidth", lambda p: fit(p.fitToWidth)),
        ("fitToHeight", lambda p: fit(p.fitToHeight)),
    ]
    for name, getter in attrs:
        bv, cv = getter(bp), getter(cp)
        if bv != cv:
            report.add("page_setup", title, name, bv, cv, "页面设置不一致")
    bm, cm = bws.page_margins, cws.page_margins
    if bm is not None or cm is not None:
        for side in ("left", "right", "top", "bottom"):
            bv = getattr(bm, side, None) if bm is not None else None
            cv = getattr(cm, side, None) if cm is not None else None
            if bv is None or cv is None:
                if bv != cv:
                    report.add("page_setup", title, f"margin.{side}", bv, cv, "页边距不一致")
            elif abs(float(bv) - float(cv)) > tol:
                report.add("page_setup", title, f"margin.{side}", bv, cv, "页边距不一致")


# ---------------------------------------------------------------------------
# 加载 / CLI
# ---------------------------------------------------------------------------
def load_workbook(path: str) -> "openpyxl.Workbook":
    """加载 xlsx（默认读存储值：公式以字符串形式参与比较）。"""
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"文件不存在: {path}")
    return openpyxl.load_workbook(str(p), data_only=False, read_only=False)


def _print_text(report: DiffReport, verbose: bool) -> None:
    print("=" * 70)
    print("XLSX DIFF REPORT")
    print(f"baseline : {report.baseline}")
    print(f"candidate: {report.candidate}")
    print(f"tolerance: {report.tol}")
    print("=" * 70)
    if report.passed:
        print("PASS ✓ 两侧完全一致（列名/合并/数字格式/页面设置/逐格 全绿）")
        return
    cats: Dict[str, int] = {}
    for d in report.diffs:
        cats[d.category] = cats.get(d.category, 0) + 1
    summary = "  ".join(f"{k}={v}" for k, v in sorted(cats.items()))
    print(f"FAIL ✗ 不一致: {len(report.diffs)} 处  [{summary}]")
    print("-" * 70)
    if verbose:
        for d in report.diffs:
            print(f"[{d.category}] sheet={d.sheet} loc={d.location}")
            print(f"    expected: {_fmt(d.expected)}")
            print(f"    actual  : {_fmt(d.actual)}")
            if d.note:
                print(f"    note    : {d.note}")
    else:
        # 非 verbose 也列出每条摘要，便于定位
        for d in report.diffs[:200]:
            print(f"[{d.category}] {d.sheet} {d.location}: {_fmt(d.expected)} -> {_fmt(d.actual)}")
        if len(report.diffs) > 200:
            print(f"... 其余 {len(report.diffs) - 200} 处省略，使用 --verbose 查看全部")
    print("-" * 70)


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="逐格跨 sheet 比对两个 xlsx（golden 回归安全网，纯 openpyxl 无头运行）",
    )
    parser.add_argument("baseline", help="基准 xlsx 路径（golden）")
    parser.add_argument("candidate", help="待比对 xlsx 路径（C# 产物等）")
    parser.add_argument("--json", dest="json_out", default=None, help="将结构化 diff 写入该 JSON 文件")
    parser.add_argument("--verbose", action="store_true", help="逐条打印每处不一致详情")
    parser.add_argument("--tol", type=float, default=1e-6, help="数值比较容差（默认 1e-6）")
    args = parser.parse_args(argv)

    try:
        base_wb = load_workbook(args.baseline)
        cand_wb = load_workbook(args.candidate)
    except Exception as exc:  # noqa: BLE001
        print(f"[ERROR] 无法读取工作簿: {exc}", file=sys.stderr)
        return 2

    report = DiffReport(baseline=args.baseline, candidate=args.candidate, tol=args.tol)
    compare_workbooks(base_wb, cand_wb, args.tol, report)

    _print_text(report, args.verbose)
    if args.json_out:
        out = Path(args.json_out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(report.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"[INFO] 结构化 diff 已写入: {out}")

    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
