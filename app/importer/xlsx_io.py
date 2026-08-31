"""openpyxl 读取封装。

源台账文件（多为财务/会计软件导出）常缺默认样式定义，openpyxl 在
load_workbook 时会打一条良性 UserWarning：
    "Workbook contains no default style, apply openpyxl's default"
数据读取完全不受影响，openpyxl 会自动套用自己的默认样式。

这里在加载期间仅屏蔽这条特定警告，避免污染运行日志/控制台；其余
UserWarning 不受影响。
"""
from __future__ import annotations

import warnings

import openpyxl


def load_workbook(path, **kwargs):
    """等价于 openpyxl.load_workbook，但屏蔽 'no default style' 良性警告。"""
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style",
            category=UserWarning,
        )
        return openpyxl.load_workbook(path, **kwargs)
