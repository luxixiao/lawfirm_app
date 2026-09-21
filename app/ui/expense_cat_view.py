"""费用类型页：6 张分类卡片 + 类内/跨类拖拽

交互口径（已与需求方确认）：
- 分类固定 6 类（报酬发放/住房公积金/保险费/汽油费/报销摊销等/公共专属费用），**说明文字可自定义填写**。
- 每张卡片列出该类下的费用类型，支持：
  · 类内拖动排序、跨卡片拖动搬运（改归类）
  · 上移/下移按钮做精确定位兜底（长列表拖拽不好定位）
  · 类内新增、改名（双击）、删除、移动到其它分类
- 顺序与归类会写入 expense_cat.sort_order，直接影响结算表/年度聘用结算表的费用列序。
"""
from __future__ import annotations

import openpyxl
from typing import Dict, List

from PySide6.QtCore import QPoint, QRect, QSize, Qt
from PySide6.QtWidgets import (
    QAbstractItemView, QComboBox, QDialog, QDialogButtonBox, QFrame, QLayout,
    QGridLayout, QHBoxLayout, QInputDialog, QFileDialog, QLabel, QLineEdit, QListWidget,
    QListWidgetItem, QMenu, QMessageBox, QScrollArea, QSizePolicy, QTextEdit, QVBoxLayout, QWidget,
)

from app.ui.widgets import CaptionLabel, PageHeader, PushButton
from app.db import get_conn
from app.engine.change_log import log_change
from app.engine import expense_cat as ec


# ---------------------------------------------------------------------------
# 流式布局（chips 自动换行；PySide6 未内置，移植 Qt 官方 FlowLayout 示例）
# ---------------------------------------------------------------------------

class FlowLayout(QLayout):
    """水平流式布局：子项从左到右排列，超出宽度自动换行。"""

    def __init__(self, parent=None, margin=0, spacing=-1) -> None:
        super().__init__(parent)
        if parent is not None:
            self.setContentsMargins(margin, margin, margin, margin)
        self.setSpacing(spacing if spacing >= 0 else 4)
        self._items = []

    def addItem(self, item) -> None:  # noqa: N802
        self._items.append(item)

    def itemAt(self, index):  # noqa: N802
        if 0 <= index < len(self._items):
            return self._items[index]
        return None

    def takeAt(self, index):  # noqa: N802
        if 0 <= index < len(self._items):
            return self._items.pop(index)
        return None

    def count(self) -> int:  # noqa: N802
        return len(self._items)

    def expandingDirections(self):  # noqa: N802
        return Qt.Orientation(0)

    def hasHeightForWidth(self) -> bool:  # noqa: N802
        return True

    def heightForWidth(self, width: int) -> int:  # noqa: N802
        return self._do_layout(QRect(0, 0, width, 0), True)

    def setGeometry(self, rect: QRect) -> None:  # noqa: N802
        super().setGeometry(rect)
        self._do_layout(rect, False)

    def sizeHint(self) -> QSize:  # noqa: N802
        return self.minimumSize()

    def minimumSize(self) -> QSize:  # noqa: N802
        size = QSize()
        for item in self._items:
            size = size.expandedTo(item.minimumSize())
        m = self.contentsMargins()
        size += QSize(2 * m.left(), 2 * m.top())
        return size

    def _do_layout(self, rect: QRect, test_only: bool) -> int:
        m = self.contentsMargins()
        x = rect.x() + m.left()
        y = rect.y() + m.top()
        line_height = 0
        spacing = self.spacing()
        for item in self._items:
            wid = item.widget()
            next_x = x + item.sizeHint().width() + spacing
            if next_x - spacing > rect.right() and line_height > 0:
                x = rect.x() + m.left()
                y = y + line_height + spacing
                next_x = x + item.sizeHint().width() + spacing
                line_height = 0
            if not test_only:
                item.setGeometry(QRect(QPoint(x, y), item.sizeHint()))
            x = next_x
            line_height = max(line_height, item.sizeHint().height())
        return y + line_height - rect.y()


# ---------------------------------------------------------------------------
# 别名 chip
# ---------------------------------------------------------------------------

class AliasChip(QFrame):
    """单个别名 chip：展示「别名」并附 × 删除按钮。

    canonical 仅用于 tooltip，删除只需 alias（落库按 alias 唯一删除）。
    """

    def __init__(self, alias: str, canonical: str, on_delete, parent=None) -> None:
        super().__init__(parent)
        self.alias = alias
        self.setObjectName("chip")
        lay = QHBoxLayout(self)
        lay.setContentsMargins(6, 1, 2, 1)
        lay.setSpacing(2)
        text = QLabel(alias)
        text.setToolTip(f"别名 → 规范类型：{canonical}")
        text.setObjectName("chipText")
        lay.addWidget(text)
        x = PushButton("×")
        x.setObjectName("chipX")
        x.setFixedSize(18, 18)
        x.clicked.connect(lambda: on_delete(alias))
        lay.addWidget(x)


# ---------------------------------------------------------------------------
# 可拖拽的类型列表（类内排序 / 跨卡片搬运）
# ---------------------------------------------------------------------------

class TypeListWidget(QListWidget):
    """费用类型列表：支持类内重排与跨卡片搬运。

    Qt 的 InternalMove 只允许同一 widget 内移动，跨 widget 必须自己接管 dropEvent：
    从 source 取走条目、按落点插入本列表，最后统一落库。
    """

    def __init__(self, category: str, view: "ExpenseCatView", parent=None) -> None:
        super().__init__(parent)
        self.category = category
        self._view = view
        self._dragging: List[str] = []
        self.setDragEnabled(True)
        self.setAcceptDrops(True)
        self.setDropIndicatorShown(True)
        self.setDragDropMode(QAbstractItemView.DragDropMode.DragDrop)
        self.setDefaultDropAction(Qt.DropAction.MoveAction)
        self.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.setSpacing(1)
        self.setUniformItemSizes(True)
        self.itemDoubleClicked.connect(self._on_rename)

    # -- 取值 -----------------------------------------------------------
    def type_names(self) -> List[str]:
        return [self.item(i).text() for i in range(self.count())]

    def set_types(self, names: List[str]) -> None:
        self.blockSignals(True)
        self.clear()
        for n in names:
            self.addItem(QListWidgetItem(n))
        self.blockSignals(False)

    # -- 拖拽 -----------------------------------------------------------
    def startDrag(self, supportedActions) -> None:  # noqa: N802
        self._dragging = [self.item(i).text() for i in range(self.count())
                          if self.item(i).isSelected()]
        super().startDrag(supportedActions)

    def _drop_row(self, event) -> int:
        idx = self.indexAt(event.position().toPoint())
        if not idx.isValid():
            return self.count()
        row = idx.row()
        if self.dropIndicatorPosition() == QAbstractItemView.DropIndicatorPosition.BelowItem:
            row += 1
        return row

    def dropEvent(self, event) -> None:  # noqa: N802
        src = event.source()
        if not isinstance(src, TypeListWidget):
            event.ignore()
            return
        names = [n for n in getattr(src, "_dragging", []) or [] if n]
        if not names:
            event.ignore()
            return
        self.apply_drop(names, src, self._drop_row(event))
        event.acceptProposedAction()

    def apply_drop(self, names: List[str], src: "TypeListWidget", drop_row: int) -> None:
        """把 names 从 src 搬到本列表的 drop_row 处（src is self 即类内重排）。"""
        if not names:
            return
        moving = set(names)

        if src is self:
            current = self.type_names()
            if drop_row < 0:
                drop_row = len(current)
            removed_before = sum(1 for t in current[:drop_row] if t in moving)
            remaining = [t for t in current if t not in moving]
            insert_at = max(0, min(drop_row - removed_before, len(remaining)))
            for k, n in enumerate(names):
                remaining.insert(insert_at + k, n)
            self.set_types(remaining)
        else:
            src.set_types([t for t in src.type_names() if t not in moving])
            current = self.type_names()
            insert_at = max(0, min(drop_row, len(current)))
            for k, n in enumerate(names):
                current.insert(insert_at + k, n)
            self.set_types(current)
        for n in names:
            self._select(n)
        self._view.persist()

    def _select(self, name: str) -> None:
        for i in range(self.count()):
            if self.item(i).text() == name:
                self.item(i).setSelected(True)

    # -- 改名 -----------------------------------------------------------
    def _on_rename(self, item: QListWidgetItem) -> None:
        old = item.text()
        new, ok = QInputDialog.getText(self, "改名费用类型", "新的费用类型名称：", text=old)
        if not ok:
            return
        new = new.strip()
        if not new or new == old:
            return
        self._view.rename_type(old, new)


# ---------------------------------------------------------------------------
# 分类卡片
# ---------------------------------------------------------------------------

class CategoryCard(QFrame):
    def __init__(self, name: str, note: str, view: "ExpenseCatView") -> None:
        super().__init__()
        self.name = name
        self._view = view
        self.setObjectName("card")
        self.setFrameShape(QFrame.Shape.NoFrame)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(10, 8, 10, 8)
        lay.setSpacing(6)

        head = QHBoxLayout()
        self.title = QLabel(name)
        self.title.setObjectName("cardTitle")
        self.count = CaptionLabel("0")
        head.addWidget(self.title)
        head.addWidget(self.count)
        head.addStretch()
        btn_note = PushButton("说明")
        btn_note.setObjectName("cardBtn")      # 与底部操作按钮一致：紧凑样式 + 不裁字
        btn_note.setFixedHeight(24)
        # 按文字宽度设最小宽：窄卡片（窗口被压窄时）也不裁字
        fm = btn_note.fontMetrics()
        btn_note.setMinimumWidth(fm.horizontalAdvance("说明") + 26)
        btn_note.clicked.connect(self._edit_note)
        head.addWidget(btn_note)
        lay.addLayout(head)

        self.note_label = CaptionLabel(note or "（点击「说明」填写本类的口径说明）")
        self.note_label.setWordWrap(True)
        if note:
            self.note_label.setToolTip(note)
        lay.addWidget(self.note_label)

        self.list = TypeListWidget(name, view)
        self.list.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        lay.addWidget(self.list, 1)

        # 别名区：按规范类型分组展示别名 chips（导入时自动归一到规范名）
        self.alias_area = QVBoxLayout()
        self.alias_area.setContentsMargins(0, 0, 0, 0)
        self.alias_area.setSpacing(4)
        self._alias_host = QWidget()
        self._alias_host.setLayout(self.alias_area)
        lay.addWidget(self._alias_host)

        bar = QHBoxLayout()
        bar.setSpacing(4)
        b_add = PushButton("新增")
        b_add.clicked.connect(self._add)
        b_up = PushButton("上移")
        b_up.setToolTip("在本类内向上移一位")
        b_up.clicked.connect(lambda: self._move(-1))
        b_down = PushButton("下移")
        b_down.setToolTip("在本类内向下移一位")
        b_down.clicked.connect(lambda: self._move(1))
        # 「移动」而非「移动到」：少一字，五个按钮一行时才放得下（窄窗口不裁字）
        self.b_move = PushButton("移动")
        self.b_move.setToolTip("移动到其它分类")
        menu = QMenu(self.b_move)
        for c in ec.CATEGORIES:
            if c != name:
                menu.addAction(c, lambda _=False, c=c: self._move_to(c))
        self.b_move.setMenu(menu)
        b_del = PushButton("删除")
        b_del.clicked.connect(self._delete)
        for b in (b_add, b_up, b_down, self.b_move, b_del):
            b.setObjectName("cardBtn")      # 紧凑样式，避免被压到 sizeHint 以下
            b.setFixedHeight(24)
            # 按文字宽度设最小宽：窄卡片（窗口被压窄时）也不会裁字。
            # 带菜单的「移动」多留箭头位；其余按 2 字文字 + 留白即可。
            fm = b.fontMetrics()
            extra = 40 if b.menu() is not None else 26
            b.setMinimumWidth(fm.horizontalAdvance(b.text()) + extra)
            bar.addWidget(b)
        bar.addStretch(1)                   # 富余空间留在右侧，按钮保持自身宽度
        lay.addLayout(bar)

    # -- 数据 -----------------------------------------------------------
    def set_types(self, names: List[str]) -> None:
        self.list.set_types(names)
        self.count.setText(f"{len(names)} 项")

    def set_note(self, note: str) -> None:
        txt = (note or "").strip()
        self.note_label.setText(txt or "（点击「说明」填写本类的口径说明）")
        self.note_label.setToolTip(txt)

    def selected(self) -> List[str]:
        return [self.list.item(i).text() for i in range(self.list.count())
                if self.list.item(i).isSelected()]

    # -- 别名 -----------------------------------------------------------
    def set_aliases(self, aliases_by_canonical: Dict[str, List[str]]) -> None:
        """按本卡片的类型分组渲染别名 chips；无别名时给操作提示。"""
        while self.alias_area.count():
            item = self.alias_area.takeAt(0)
            w = item.widget()
            if w is not None:
                w.deleteLater()
        types = self.list.type_names()
        any_alias = False
        for t in types:
            al = aliases_by_canonical.get(t, [])
            if not al:
                continue
            any_alias = True
            row = QWidget()
            rlay = QHBoxLayout(row)
            rlay.setContentsMargins(0, 0, 0, 0)
            rlay.setSpacing(4)
            lbl = QLabel(t)
            lbl.setObjectName("chipType")
            rlay.addWidget(lbl)
            flow_host = QWidget()
            FlowLayout(flow_host, margin=0, spacing=4)
            for a in al:
                flow_host.layout().addWidget(AliasChip(a, t, self._view.delete_alias))
            rlay.addWidget(flow_host, 1)
            self.alias_area.addWidget(row)
        if not any_alias:
            hint = CaptionLabel(
                "别名：同义写法（如「公积金」→「住房公积金」），导入时自动归一到规范名；"
                "点下方「＋别名」添加。")
            hint.setWordWrap(True)
            self.alias_area.addWidget(hint)
        b_add = PushButton("＋别名")
        b_add.setObjectName("cardBtn")
        b_add.setFixedHeight(24)
        b_add.clicked.connect(self._add_alias)
        self.alias_area.addWidget(b_add)

    def _add_alias(self) -> None:
        types = self.list.type_names()
        if not types:
            QMessageBox.information(self, "提示", "请先在卡片中新增一个规范类型，再为它添加别名")
            return
        dlg = AliasAddDialog(types, self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            canonical, alias = dlg.value()
            if alias:
                self._view.add_alias(canonical, alias)

    # -- 操作 -----------------------------------------------------------
    def _edit_note(self) -> None:
        dlg = NoteDialog(self.name, ec.get_category_note(self.name), self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            note = dlg.value()
            ec.set_category_note(self.name, note)
            self.set_note(note)

    def _add(self) -> None:
        name, ok = QInputDialog.getText(self, "新增费用类型",
                                        f"在「{self.name}」下新增费用类型：")
        if not ok:
            return
        self._view.add_type(name, self.name)

    def _move(self, direction: int) -> None:
        sel = self.selected()
        if not sel:
            QMessageBox.information(self, "提示", "请先选中要移动的费用类型")
            return
        self._view.move_in_category(sel[0], direction)

    def _move_to(self, category: str) -> None:
        sel = self.selected()
        if not sel:
            QMessageBox.information(self, "提示", "请先选中要移动的费用类型")
            return
        self._view.move_to_category(sel, category)

    def _delete(self) -> None:
        sel = self.selected()
        if not sel:
            QMessageBox.information(self, "提示", "请先选中要删除的费用类型")
            return
        self._view.delete_types(sel, self.name)


class NoteDialog(QDialog):
    """分类说明编辑（多行文本）。"""

    def __init__(self, category: str, note: str, parent=None) -> None:
        super().__init__(parent)
        self.setWindowTitle(f"说明 — {category}")
        self.resize(420, 220)
        lay = QVBoxLayout(self)
        lay.addWidget(CaptionLabel(f"「{category}」这一类的口径说明，可自定义填写（留空即可）："))
        self.edit = QTextEdit()
        self.edit.setPlainText(note or "")
        lay.addWidget(self.edit, 1)
        box = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok
                               | QDialogButtonBox.StandardButton.Cancel)
        box.button(QDialogButtonBox.StandardButton.Ok).setText("确定")
        box.button(QDialogButtonBox.StandardButton.Cancel).setText("取消")
        box.accepted.connect(self.accept)
        box.rejected.connect(self.reject)
        lay.addWidget(box)

    def value(self) -> str:
        return self.edit.toPlainText().strip()


class AliasAddDialog(QDialog):
    """新增别名：选择规范类型 + 输入同义写法。"""

    def __init__(self, types: List[str], parent=None) -> None:
        super().__init__(parent)
        self.setWindowTitle("新增别名")
        self.resize(420, 180)
        lay = QVBoxLayout(self)
        lay.addWidget(CaptionLabel("为以下规范类型添加一个同义别名（导入时自动归一到它）："))
        self.combo = QComboBox()
        self.combo.addItems(types)
        lay.addWidget(self.combo)
        lay.addWidget(CaptionLabel("别名（费用台账里可能出现的写法，如「公积金」）："))
        self.edit = QLineEdit()
        lay.addWidget(self.edit, 1)
        box = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok
                               | QDialogButtonBox.StandardButton.Cancel)
        box.button(QDialogButtonBox.StandardButton.Ok).setText("确定")
        box.button(QDialogButtonBox.StandardButton.Cancel).setText("取消")
        box.accepted.connect(self.accept)
        box.rejected.connect(self.reject)
        lay.addWidget(box)

    def value(self) -> tuple:
        return (self.combo.currentText(), self.edit.text().strip())


# ---------------------------------------------------------------------------
# 页面
# ---------------------------------------------------------------------------

class ExpenseCatView(QWidget):
    def __init__(self) -> None:
        super().__init__()
        self._cards: Dict[str, CategoryCard] = {}

        lay = QVBoxLayout(self)
        lay.setContentsMargins(28, 24, 28, 20)
        lay.setSpacing(12)

        lay.addWidget(PageHeader(
            "费用类型",
            "按 6 类分组维护费用类型。卡片内拖动可调整顺序，拖到别的卡片即改归类；"
            "双击类型名可改名。顺序与归类决定结算表的费用列序。",
        ))

        self.scroll = QScrollArea()
        self.scroll.setWidgetResizable(True)
        self.scroll.setFrameShape(QFrame.Shape.NoFrame)
        self.body = QWidget()
        self.grid = QGridLayout(self.body)
        self.grid.setContentsMargins(0, 0, 8, 0)
        self.grid.setSpacing(12)
        # 尊重各卡片最小宽 300：窗口被压窄到放不下 3 列时，滚动而非压缩卡片
        self.grid.setSizeConstraint(QLayout.SizeConstraint.SetMinimumSize)
        self.scroll.setWidget(self.body)
        lay.addWidget(self.scroll, 1)

        foot = QHBoxLayout()
        self.hint = CaptionLabel("")
        foot.addWidget(self.hint)
        foot.addStretch()
        btn_export = PushButton("导出")
        btn_export.setToolTip("将所有费用类型（分类/说明/归类/顺序）导出为 Excel")
        btn_export.clicked.connect(self._export)
        foot.addWidget(btn_export)
        btn_import = PushButton("导入")
        btn_import.setToolTip("从 Excel 整表替换所有费用类型配置")
        btn_import.clicked.connect(self._import)
        foot.addWidget(btn_import)
        btn_reset = PushButton("刷新")
        btn_reset.clicked.connect(self.refresh)
        foot.addWidget(btn_reset)
        lay.addLayout(foot)

        self.refresh()

    def showEvent(self, event) -> None:  # noqa: N802
        super().showEvent(event)
        self.refresh()

    # -- 渲染 -----------------------------------------------------------
    def refresh(self) -> None:
        cats = ec.list_categories()
        # 首次/分类数量变化时重建卡片
        if {c["name"] for c in cats} != set(self._cards):
            while self.grid.count():
                item = self.grid.takeAt(0)
                w = item.widget()
                if w is not None:
                    w.deleteLater()
            self._cards = {}
            for i, c in enumerate(cats):
                card = CategoryCard(c["name"], c["note"], self)
                # 300 是按钮行（新增/上移/下移/移动/删除）不裁字的最小宽度
                card.setMinimumWidth(300)
                card.setMinimumHeight(240)
                self._cards[c["name"]] = card
                self.grid.addWidget(card, i // 3, i % 3)
            # 列可拉伸：窗口宽时卡片变宽，窄时不至于被压到最小宽度以下
            for col in range(3):
                self.grid.setColumnStretch(col, 1)
            self.grid.setRowStretch(len(cats) // 3 + 1, 1)
        by_cat = ec.types_by_category()
        aliases_by_canonical = ec.aliases_by_canonical()
        for c in cats:
            card = self._cards.get(c["name"])
            if card is None:
                continue
            card.set_types(by_cat.get(c["name"], []))
            card.set_note(c["note"])
            card.set_aliases(aliases_by_canonical)
        total = sum(len(v) for v in by_cat.values())
        self.hint.setText(f"共 {total} 个费用类型，{len(cats)} 个分类")

    # -- 配置导入 / 导出（Excel） --------------------------------------
    def _export(self) -> None:
        path, _ = QFileDialog.getSaveFileName(
            self, "导出费用类型", "费用类型配置.xlsx", "Excel 文件 (*.xlsx)")
        if not path:
            return
        try:
            _export_excel(path, ec.export_all())
        except Exception as e:  # noqa: BLE001
            QMessageBox.critical(self, "导出失败", str(e))
            return
        QMessageBox.information(self, "已导出", f"已导出费用类型配置到：\n{path}")

    def _import(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, "导入费用类型", "", "Excel 文件 (*.xlsx *.xlsm)")
        if not path:
            return
        try:
            data = _import_excel(path)
        except Exception as e:  # noqa: BLE001
            QMessageBox.critical(self, "导入失败", str(e))
            return
        if not data["types"]:
            QMessageBox.warning(self, "导入失败", "文件中没有可用的费用类型数据")
            return
        if QMessageBox.question(
                self, "确认导入",
                f"导入将【整表替换】当前全部费用类型配置（共 {len(data['types'])} 个类型）。\n"
                "建议先导出一份备份。确定继续？") != QMessageBox.StandardButton.Yes:
            return
        try:
            ec.import_all(data)
        except Exception as e:  # noqa: BLE001
            QMessageBox.critical(self, "导入失败", str(e))
            return
        self.refresh()
        QMessageBox.information(self, "已导入", "费用类型配置已更新。")

    # -- 落库 -----------------------------------------------------------
    def _layout_now(self) -> Dict[str, List[str]]:
        return {name: card.list.type_names() for name, card in self._cards.items()}

    def persist(self) -> None:
        """拖拽/排序后统一落库（归类 + 全局顺序）。"""
        ec.save_layout(self._layout_now())
        for name, card in self._cards.items():
            card.count.setText(f"{card.list.count()} 项")

    def add_type(self, name: str, category: str) -> None:
        name = (name or "").strip()
        if not name:
            return
        try:
            ec.add_type(name, category)
        except ec.ExpenseCatError as exc:
            QMessageBox.warning(self, "无法新增", str(exc))
            return
        conn = get_conn()
        try:
            log_change(conn, "expense_cat", name, "expense_type", "", name, "费用类型维护")
            conn.commit()
        finally:
            conn.close()
        self.refresh()

    def rename_type(self, old: str, new: str) -> None:
        try:
            ec.rename_type(old, new)
        except ec.ExpenseCatError as exc:
            QMessageBox.warning(self, "无法改名", str(exc))
            return
        conn = get_conn()
        try:
            log_change(conn, "expense_cat", old, "expense_type", old, new, "费用类型维护")
            conn.commit()
        finally:
            conn.close()
        self.refresh()

    def delete_types(self, names: List[str], category: str) -> None:
        if not names:
            return
        if len(names) == 1:
            n = ec.type_reference_count(names[0])
            tip = (f"确定删除费用类型「{names[0]}」？")
            if n:
                tip += (f"\n\n该类型在费用台账中有 {n} 条记录。"
                        f"\n删除后这些记录仍保留，但结算表不会再单列该类型。")
        else:
            n = sum(ec.type_reference_count(x) for x in names)
            tip = f"确定删除选中的 {len(names)} 个费用类型？"
            if n:
                tip += f"\n\n这些类型在费用台账中共 {n} 条记录。"
        if QMessageBox.question(self, "删除费用类型", tip) != QMessageBox.StandardButton.Yes:
            return
        for name in names:
            ec.delete_type(name)
            conn = get_conn()
            try:
                log_change(conn, "expense_cat", name, "expense_type", name, "", "费用类型维护")
                conn.commit()
            finally:
                conn.close()
        self.refresh()

    def add_alias(self, canonical: str, alias: str) -> None:
        try:
            ec.add_alias(canonical, alias)
        except ec.ExpenseCatError as exc:
            QMessageBox.warning(self, "无法新增别名", str(exc))
            return
        conn = get_conn()
        try:
            log_change(conn, "expense_type_alias", alias, "alias", "", alias,
                       f"费用类型别名→{canonical}")
            conn.commit()
        finally:
            conn.close()
        self.refresh()

    def delete_alias(self, alias: str) -> None:
        if QMessageBox.question(
                self, "删除别名",
                f"确定删除别名「{alias}」？\n删除后费用台账里若再出现该写法，将重新按「未知类型」报错。"
        ) != QMessageBox.StandardButton.Yes:
            return
        try:
            ec.delete_alias(alias)
        except ec.ExpenseCatError as exc:
            QMessageBox.warning(self, "无法删除别名", str(exc))
            return
        conn = get_conn()
        try:
            log_change(conn, "expense_type_alias", alias, "alias", alias, "",
                       "费用类型别名删除")
            conn.commit()
        finally:
            conn.close()
        self.refresh()

    def move_in_category(self, name: str, direction: int) -> None:
        if not ec.move_in_category(name, direction):
            return
        self.refresh()
        self._reselect(name)

    def move_to_category(self, names: List[str], category: str) -> None:
        for name in names:
            old = self._category_of(name)
            if old == category:
                continue
            ec.set_category(name, category)
            conn = get_conn()
            try:
                log_change(conn, "expense_cat", name, "category", old, category, "费用类型维护")
                conn.commit()
            finally:
                conn.close()
        self.refresh()
        self._reselect(names[-1])

    def _category_of(self, name: str) -> str:
        m = ec.get_map()
        cat = m.get(name, ec.FALLBACK_CATEGORY)
        return cat if cat in ec.CATEGORIES else ec.FALLBACK_CATEGORY

    def _reselect(self, name: str) -> None:
        for card in self._cards.values():
            for i in range(card.list.count()):
                it = card.list.item(i)
                it.setSelected(it.text() == name)


# ---------------------------------------------------------------------------
# Excel 配置导入 / 导出（费用类型整表备份迁移）
# ---------------------------------------------------------------------------

def _export_excel(path: str, data: Dict) -> None:
    """把费用类型配置写成三 sheet Excel：分类 / 类型 / 别名。"""
    wb = openpyxl.Workbook()
    ws_cat = wb.active
    ws_cat.title = "分类"
    ws_cat.append(["分类", "说明", "顺序"])
    for c in data["categories"]:
        ws_cat.append([c["name"], c.get("note", ""), c.get("sort_order", 0)])
    ws_type = wb.create_sheet("类型")
    ws_type.append(["类型", "分类", "顺序"])
    for t in data["types"]:
        ws_type.append([t["expense_type"], t["category"], t.get("sort_order", 0)])
    ws_alias = wb.create_sheet("别名")
    ws_alias.append(["别名", "规范类型"])
    for a in data.get("aliases", []):
        ws_alias.append([a["alias"], a["canonical"]])
    wb.save(path)


def _import_excel(path: str) -> Dict:
    """读取 Excel，还原为 {categories, types, aliases}（兼容旧格式无别名 sheet）。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    cats, types, aliases = [], [], []
    if "分类" in wb.sheetnames:
        for row in wb["分类"].iter_rows(min_row=2, values_only=True):
            if row and row[0] not in (None, ""):
                cats.append({
                    "name": str(row[0]).strip(),
                    "note": str(row[1]) if len(row) > 1 and row[1] is not None else "",
                    "sort_order": int(row[2]) if len(row) > 2 and row[2] is not None else 0,
                })
    if "类型" in wb.sheetnames:
        for row in wb["类型"].iter_rows(min_row=2, values_only=True):
            if row and row[0] not in (None, ""):
                types.append({
                    "expense_type": str(row[0]).strip(),
                    "category": str(row[1]).strip() if len(row) > 1 and row[1] else ec.FALLBACK_CATEGORY,
                    "sort_order": int(row[2]) if len(row) > 2 and row[2] is not None else 0,
                })
    if "别名" in wb.sheetnames:
        for row in wb["别名"].iter_rows(min_row=2, values_only=True):
            if row and row[0] not in (None, ""):
                aliases.append({
                    "alias": str(row[0]).strip(),
                    "canonical": str(row[1]).strip() if len(row) > 1 and row[1] is not None else "",
                })
    if not cats and not types and wb.sheetnames:
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] not in (None, "") and len(row) >= 2 and row[1] not in (None, ""):
                cats.append({"name": str(row[0]).strip(),
                             "note": str(row[3]) if len(row) > 3 and row[3] is not None else "",
                             "sort_order": 0})
                types.append({"expense_type": str(row[1]).strip(),
                              "category": str(row[0]).strip(),
                              "sort_order": int(row[2]) if len(row) > 2 and row[2] is not None else 0})
    return {"categories": cats, "types": types, "aliases": aliases}
