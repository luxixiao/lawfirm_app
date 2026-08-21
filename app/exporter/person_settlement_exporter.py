"""个人结算总表 Excel 导出器（按模板样式，每人一份文件）"""
from __future__ import annotations

from pathlib import Path
from typing import Dict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.engine.person_settlement import MONTHS, build_settlement

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")


def _row_values(st: Dict, months: list, key: str) -> list:
    """某 key 的 1~12 月值 + 合计"""
    vals = [round(st["months"][m][key], 2) for m in months]
    return vals + [round(sum(vals), 2)]


def _subtotal(st: Dict, months: list, keys: list) -> list:
    vals = [round(sum(st["months"][m][k] for k in keys), 2) for m in months]
    return vals + [round(sum(vals), 2)]


def export_one(st: Dict, person: str, path: str | Path, year: int) -> Path:
    """导出单个员工的个人结算总表（模板样式）"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "个人结算总表"

    # 标题
    ws.merge_cells("A1:O1")
    c = ws["A1"]
    c.value = f"个人结算总表（{year}年）　姓名：{person}　类型：{st['staff_type']}"
    c.font = Font(size=14, bold=True)
    c.alignment = CENTER
    ws.row_dimensions[1].height = 28

    # 表头
    headers = ["序号", "项目", "1月", "2月", "3月", "4月", "5月", "6月",
               "7月", "8月", "9月", "10月", "11月", "12月", "合计"]
    for j, h in enumerate(headers, 1):
        cell = ws.cell(2, j, h)
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = BORDER

    def put(row_idx: int, seq: str, name: str, vals: list, bold: bool = False, seq_only: bool = False):
        ws.cell(row_idx, 1, seq if not seq_only else "").alignment = CENTER
        nm = ws.cell(row_idx, 2, name)
        nm.alignment = Alignment(horizontal="left", vertical="center")
        if bold:
            nm.font = BOLD
        for j, v in enumerate(vals, 3):
            cell = ws.cell(row_idx, j, None if v is None else round(v, 2))
            cell.number_format = "#,##0.00"
            cell.alignment = RIGHT
            cell.border = BORDER
        ws.cell(row_idx, 1).border = BORDER
        ws.cell(row_idx, 2).border = BORDER
        if bold:
            for j in range(3, 16):
                ws.cell(row_idx, j).font = BOLD

    r = 3
    months = list(range(1, 13))
    m = st["months"]

    # 一、上年结余结转（留空）
    put(r, "一", "上年结余结转", [None] * 13); r += 1

    # 二、本月收款金额（小计）
    rec_keys = ["rec_open_cur", "rec_cur_year", "rec_prev_year", "rec_refund_cur", "rec_refund_prev"]
    put(r, "二", "本月收款金额", _subtotal(st, months, rec_keys), bold=True); r += 1
    put(r, "1", "本月开收", _row_values(st, months, "rec_open_cur")); r += 1
    put(r, "2", "收本年", _row_values(st, months, "rec_cur_year")); r += 1
    put(r, "3", "收上年", _row_values(st, months, "rec_prev_year")); r += 1
    put(r, "4", "退本年", _row_values(st, months, "rec_refund_cur")); r += 1
    put(r, "5", "退上年", _row_values(st, months, "rec_refund_prev")); r += 1

    # 三、本月开具发票金额
    inv_keys = ["inv_open_received", "inv_open_uncollected", "inv_red_cur", "inv_red_prev"]
    put(r, "三", "本月开具发票金额", _subtotal(st, months, inv_keys), bold=True); r += 1
    put(r, "1", "本月开收", _row_values(st, months, "inv_open_received")); r += 1
    put(r, "2", "本月未收", _row_values(st, months, "inv_open_uncollected")); r += 1
    put(r, "3", "红冲本年", _row_values(st, months, "inv_red_cur")); r += 1
    put(r, "4", "红冲上年", _row_values(st, months, "inv_red_prev")); r += 1

    # 四、未收款金额（各月=本月未收；合计=本年累计未收）
    uncollected_vals = [round(st["uncollected_month"][mo], 2) for mo in months] + [round(st["uncollected_total"], 2)]
    put(r, "四", "未收款金额", uncollected_vals, bold=True); r += 1

    # 五、业务收入
    put(r, "五", "业务收入", _row_values(st, months, "income"), bold=True); r += 1

    # 六、减：分成报酬及费用
    exp = st["expenses"]
    exp_sub = [round(sum(exp.get(t, {}).get(mo, 0.0) for t in exp), 2) for mo in months]
    exp_sub.append(round(sum(exp_sub), 2))
    put(r, "六", "减：分成报酬及费用", exp_sub, bold=True); r += 1
    for i, etype in enumerate(sorted(exp.keys()), 1):
        vals = [round(exp[etype].get(mo, 0.0), 2) for mo in months]
        vals.append(round(sum(vals), 2))
        put(r, str(i), etype, vals); r += 1

    # 列宽
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 22
    for j in range(3, 16):
        ws.column_dimensions[get_column_letter(j)].width = 11

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def export_all(out_dir: str | Path, year: int) -> list:
    """为全部员工生成个人结算总表，返回生成的文件路径列表"""
    data = build_settlement(year)
    out_dir = Path(out_dir)
    files = []
    for person, st in data.items():
        safe = person.replace("/", "_").replace("\\", "_").strip() or "未命名"
        path = export_one(st, person, out_dir / f"个人结算总表_{safe}.xlsx", year)
        files.append(path)
    return files
