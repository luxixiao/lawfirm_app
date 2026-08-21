"""月度结算表导出器：仿模板生成一个多 sheet Excel（每人一个 sheet + 公共费用）

结构（仿「2025年1月结算表.xlsx」）：
- 列：序号 | 项目 | 本期 | 本年累计 | 备注
- 行：一 上年结余结转 / 二 本月收款金额（1.本月开票本月收款 2.收回以前应收款）
     / 三 本月开具发票金额（1.本月开票已收款 2.本月开票未收款）
     / 四 未收款金额（本期=本月未收；累计=本年累计未收）
     / 五 业务收入 / 六 减：分成报酬及费用（逐类型）
     / 七 减：税金及会费（增值税、城建及附加=开票净额÷1.06×6%×1.12；抵扣进项留空）
     / 结余 = 五 - 六 - 七
备注：收回以前应收款>0 时自动生成
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from app.engine.person_settlement import build_settlement

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")

CN_NUM = ["一", "二", "三", "四", "五", "六", "七"]


def _tax(amount: float) -> float:
    """增值税、城建及附加 = 开票净额 ÷ 1.06 × 6% × 1.12"""
    return round(amount / 1.06 * 0.06 * 1.12, 2)


def _chinese_title(month: int) -> str:
    cn = "一二三四五六七八九十"
    if month <= 10:
        return cn[month - 1] + "月"
    return "十" + (cn[month - 11] if month > 11 else "") + "月"


def _month_total(st: Dict, month: int, keys) -> float:
    return round(sum(st["months"][month][k] for k in keys), 2)


def report_note_rec(st: Dict, month: int) -> str:
    """收款类备注（写在「二、本月收款金额」行）"""
    m = st["months"]
    notes = []
    if m[month]["rec_cur_year"] > 0.01:
        notes.append(f"本月收回本年应收款{m[month]['rec_cur_year']:,.2f}元")
    if m[month]["rec_prev_year"] > 0.01:
        notes.append(f"本月收回上年应收款{m[month]['rec_prev_year']:,.2f}元")
    if m[month]["rec_refund_cur"] < -0.01:
        notes.append(f"本月退本年应收款{abs(m[month]['rec_refund_cur']):,.2f}元")
    if m[month]["rec_refund_prev"] < -0.01:
        notes.append(f"本月退上年应收款{abs(m[month]['rec_refund_prev']):,.2f}元")
    cum_prev_year = sum(m[mo]["rec_prev_year"] for mo in range(1, month + 1))
    cum_refund_prev = sum(m[mo]["rec_refund_prev"] for mo in range(1, month + 1))
    if cum_prev_year > 0.01:
        notes.append(f"本年累计收回上年应收款{cum_prev_year:,.2f}元")
    if cum_refund_prev < -0.01:
        notes.append(f"本年累计退上年应收款{abs(cum_refund_prev):,.2f}元")
    return "\n".join(notes)


def report_note_inv(st: Dict, month: int) -> str:
    """开票类备注（写在「三、本月开具发票金额」行）"""
    m = st["months"]
    notes = []
    if m[month]["inv_red_cur"] < -0.01:
        notes.append(f"本月红冲本年{abs(m[month]['inv_red_cur']):,.2f}元")
    if m[month]["inv_red_prev"] < -0.01:
        notes.append(f"本月红冲上年{abs(m[month]['inv_red_prev']):,.2f}元")
    cum_red_prev = sum(m[mo]["inv_red_prev"] for mo in range(1, month + 1))
    if cum_red_prev < -0.01:
        notes.append(f"本年累计红冲上年{abs(cum_red_prev):,.2f}元")
    return "\n".join(notes)


def build_report_rows(st: Dict, year: int, month: int) -> list:
    """结算表行数据：[(seq, name, cur, total, bold), ...]（预览与导出共用）"""
    m = st["months"]
    months = list(range(1, 13))
    cum = lambda vals12: round(sum(vals12[:month]), 2)  # noqa: E731
    rows = []

    def row(seq, name, cur, total, bold=False):
        rows.append([seq, name, None if cur is None else round(cur, 2),
                     None if total is None else round(total, 2), bold])

    # 一、上年结余结转
    row("一", "上年结余结转", None, None)
    # 二、本月收款金额
    rec_keys = ["rec_open_cur", "rec_cur_year", "rec_prev_year", "rec_refund_cur", "rec_refund_prev"]
    cur2 = _month_total(st, month, rec_keys)
    tot2 = round(sum(_month_total(st, mo, rec_keys) for mo in months[:month]), 2)
    row("二", "本月收款金额", cur2, tot2, bold=True)
    row("1", "本月开票本月收款", m[month]["rec_open_cur"],
        cum([m[mo]["rec_open_cur"] for mo in months]))
    row("2", "收回以前应收款", round(m[month]["rec_cur_year"] + m[month]["rec_prev_year"], 2),
        cum([m[mo]["rec_cur_year"] + m[mo]["rec_prev_year"] for mo in months]))
    # 三、本月开具发票金额
    cur3 = m[month]["inv_total"]
    tot3 = round(sum(m[mo]["inv_total"] for mo in months[:month]), 2)
    row("三", "本月开具发票金额", cur3, tot3, bold=True)
    row("1", "本月开票已收款", m[month]["inv_open_received"],
        cum([m[mo]["inv_open_received"] for mo in months]))
    row("2", "本月开票未收款", m[month]["inv_open_uncollected"],
        cum([m[mo]["inv_open_uncollected"] for mo in months]))
    # 四、未收款金额
    row("四", "未收款金额", st["uncollected_month"][month], st["uncollected_total"], bold=True)
    # 五、业务收入
    row("五", "业务收入", m[month]["income"], cum([m[mo]["income"] for mo in months]), bold=True)
    # 六、减：分成报酬及费用（排除折旧/银行结息/城建税及附加——模板口径）
    _EXCLUDE = ("折旧", "银行结息", "城建税", "教育附加")
    exp = {t: v for t, v in st["expenses"].items() if not any(k in t for k in _EXCLUDE)}
    cur6 = round(sum(exp.get(t, {}).get(month, 0.0) for t in exp), 2)
    tot6 = round(sum(exp.get(t, {}).get(mo, 0.0) for t in exp for mo in months[:month]), 2)
    row("六", "减：分成报酬及费用", cur6, tot6, bold=True)
    _EXP_ORDER = ["分成报酬", "合办人员报酬", "社保", "刷卡汽油费", "发票报销", "停车费",
                  "高温费", "旅游费", "行政工资", "行政年终奖", "实习工资", "公积金"]
    exp_types = sorted(exp.keys(), key=lambda t: (_EXP_ORDER.index(t) if t in _EXP_ORDER else 99, t))
    for i, etype in enumerate(exp_types, 1):
        row(str(i), etype, exp[etype].get(month, 0.0), cum([exp[etype].get(mo, 0.0) for mo in months]))
    # 七、减：税金及会费
    tax_cur = _tax(max(cur3, 0.0))
    tax_tot = _tax(max(tot3, 0.0))
    row("七", "减：税金及会费", round(tax_cur, 2), round(tax_tot, 2), bold=True)
    row("1", "增值税、城建及附加", tax_cur, tax_tot)
    row("2", "抵扣进项", None, None)
    # 结余
    row("", "结余", round(cur5 := m[month]["income"] - cur6 - tax_cur, 2),
        round(tot5 := cum([m[mo]["income"] for mo in months]) - tot6 - tax_tot, 2), bold=True)
    return rows


def _write_sheet(ws, person: str, st: Dict, year: int, month: int) -> None:
    m = st["months"]
    months = list(range(1, 13))

    # 标题
    ws.merge_cells("A1:E1")
    ws["A1"] = "浙江震天律师事务所"
    ws["A1"].font = Font(size=13, bold=True)
    ws["A1"].alignment = CENTER
    ws["B2"] = f"姓名：{person}"
    ws["C2"] = f"二〇{'{:02d}'.format(year % 100)}年{_chinese_title(month)}结算表"

    # 表头
    for j, h in enumerate(["序号", "项目", "本期", "本年累计", "备注"], 1):
        c = ws.cell(3, j, h)
        c.font = BOLD
        c.alignment = CENTER
        c.border = BORDER

    for seq, name, cur, total, bold in build_report_rows(st, year, month):
        r = ws.max_row + 1
        ws.cell(r, 1, seq).alignment = CENTER
        ws.cell(r, 2, name).alignment = Alignment(horizontal="left", vertical="center")
        c3 = ws.cell(r, 3, cur)
        c4 = ws.cell(r, 4, total)
        for c in (c3, c4):
            if c.value is not None:
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right", vertical="center")
        for j in range(1, 6):
            ws.cell(r, j).border = BORDER
        if bold:
            ws.cell(r, 2).font = BOLD
            for j in (3, 4):
                ws.cell(r, j).font = BOLD
        # 备注按类别写在对应行（二收款 / 三开票）
        if name == "本月收款金额":
            note = report_note_rec(st, month)
            if note:
                c = ws.cell(r, 5, note)
                c.alignment = Alignment(wrap_text=True, vertical="center")
                ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 50)
        elif name == "本月开具发票金额":
            note = report_note_inv(st, month)
            if note:
                c = ws.cell(r, 5, note)
                c.alignment = Alignment(wrap_text=True, vertical="center")
                ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 50)

    # ---- 自动列宽（按内容，汉字计 2 字符）----
    def _char_w(text: str) -> int:
        return sum(2 if ord(ch) > 127 else 1 for ch in str(text or ""))

    # A 列序号固定窄宽；E 列备注保底 28（无备注时也美观）
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["E"].width = 28
    for col_idx in (2, 3, 4):
        maxw = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is None:
                    continue
                lines = str(cell.value).split("\n")
                maxw = max(maxw, max(_char_w(line) for line in lines))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(maxw + 2, 48)
    # E 列备注若内容更长则扩展
    maxw_e = 0
    for row in ws.iter_rows(min_col=5, max_col=5):
        for cell in row:
            if cell.value is None:
                continue
            maxw_e = max(maxw_e, max(_char_w(line) for line in str(cell.value).split("\n")))
    ws.column_dimensions["E"].width = min(max(28, maxw_e + 2), 60)

    # ---- 打印设置：每个 sheet 打印在一页 ----
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3)


def export_report(out_path: str | Path, year: int, month: int, persons: list[str]) -> Path:
    """生成月度结算表（多 sheet），persons 为要导出的员工名单（含"公共费用"）"""
    data = build_settlement(year)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for person in persons:
        # 模板 sheet 名为"公共费用"，引擎中费用经办人为"公共"
        st = data.get(person)
        if st is None and person == "公共费用":
            st = data.get("公共")
        if st is None:
            st = {
                "staff_type": "其他",
                "months": {mo: {k: 0.0 for k in (
                    "rec_open_cur", "rec_cur_year", "rec_prev_year", "rec_refund_cur",
                    "rec_refund_prev", "inv_open_received", "inv_open_uncollected",
                    "inv_red_cur", "inv_red_prev", "inv_total", "income")} for mo in range(1, 13)},
                "uncollected_month": {mo: 0.0 for mo in range(1, 13)},
                "uncollected_total": 0.0,
                "expenses": {},
            }
        ws = wb.create_sheet(title=person)
        _write_sheet(ws, person, st, year, month)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
