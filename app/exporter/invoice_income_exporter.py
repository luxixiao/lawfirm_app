"""2025年度开票收入表 导出器

仿模板「2025年度开票收入.xls」：
- 主表 sheets：YYYYMM（每月一个，从新到旧）
  列：序号 | 姓名 | 收入开票金额(本月/累计) | 未收款(期末未收) | 收款金额(本月开票本月收回/本月收回以前应收款/本月合计收款)
- 未收款明细 sheets：未收款明细YYYY年MM月（当年截止月）+ 未收款明细YYYY年12月（历史年度，倒序）
  列：姓名 | 发票开具月份 | 对方抬头 | 未收款金额 | 合计(按人)
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from app.db import get_conn
from app.engine.person_settlement import build_settlement
from app.engine.split import allocate_receipt

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def _main_persons(conn, year: int, month: int) -> list:
    """主表名单：合伙/聘用/兼职（不过滤 is_active，见「停用」功能已取消）；入职月份晚于当前月排除"""
    rows = conn.execute(
        "SELECT name, hire_month FROM staff "
        "WHERE staff_type IN ('合伙','聘用','兼职') ORDER BY name"
    ).fetchall()
    cur = f"{year}-{month:02d}"
    out = []
    for r in rows:
        hm = (r["hire_month"] or "").strip()
        if hm and hm > cur:
            continue
        out.append(r["name"])
    return out


def _build_main_rows(data: Dict, year: int, month: int):
    """主表行数据（预览/导出共用）

    - F 本月开票本月收回 = 引擎⑥（保持引擎现状；模板对红冲上年扣减无稳定规则，不模拟）
    - G 本月收回以前应收款 = ②+③+退款（含负值，与模板一致）
    - H = F + G
    rows: [[name, 收入本月, 收入累计, 期末未收, 本月开票已收, 收回以前, 合计收款], ...]
    totals: 6 个数值列合计
    """
    rows = []
    for name in sorted(data.keys()):
        if name == "公共":
            continue
        st = data.get(name)
        if st is None:
            rows.append([name, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0])
            continue
        m = st["months"][month]
        inv_cur = m["inv_total"]
        inv_tot = round(sum(st["months"][mo]["inv_total"] for mo in range(1, month + 1)), 2)
        uncol = st["uncollected_total"]
        open_recv = m["inv_open_received"]
        prev_recv = round(m["rec_cur_year"] + m["rec_prev_year"]
                          + m["rec_refund_cur"] + m["rec_refund_prev"], 2)
        total_recv = round(open_recv + prev_recv, 2)
        rows.append([name, inv_cur, inv_tot, uncol, open_recv, prev_recv, total_recv])
    totals = [round(sum(r[j] for r in rows), 2) for j in range(1, 7)]
    return rows, totals


def build_main_preview(year: int, month: int):
    """主表预览/导出共用：返回 (persons, rows, totals)"""
    data = build_settlement(year)
    conn = get_conn()
    try:
        persons = _main_persons(conn, year, month)
    finally:
        conn.close()
    rows, totals = _build_main_rows(data, year, month)
    return persons, rows, totals


def _write_main_sheet(ws, year: int, month: int, data: Dict) -> None:
    ws.merge_cells("A1:H1")
    ws["A1"] = "浙江震天律师事务所"
    ws["A1"].font = Font(size=13, bold=True)
    ws["A1"].alignment = CENTER
    ws.merge_cells("A2:H2")
    ws["A2"] = f"{year}年{month}月律师收费情况表（开票收入）"
    ws["A2"].alignment = CENTER

    # 表头（双行）
    ws.cell(3, 1, "序号").alignment = CENTER
    ws.cell(3, 2, "姓名").alignment = CENTER
    ws.merge_cells("C3:D3")
    ws.cell(3, 3, "收入开票金额").alignment = CENTER
    ws.cell(3, 5, "未收款").alignment = CENTER
    ws.merge_cells("F3:H3")
    ws.cell(3, 6, "收款金额").alignment = CENTER
    ws.cell(4, 3, "本月数")
    ws.cell(4, 4, "累计数")
    ws.cell(4, 5, "期末未收金额")
    ws.cell(4, 6, "本月开票本月收回")
    ws.cell(4, 7, "本月收回以前应收款")
    ws.cell(4, 8, "本月合计收款")
    for r in (3, 4):
        for c in range(1, 9):
            cell = ws.cell(r, c)
            cell.font = BOLD
            cell.alignment = CENTER
            cell.border = BORDER

    rows, totals = _build_main_rows(data, year, month)
    r = 5
    for i, row in enumerate(rows, 1):
        ws.cell(r, 1, i).alignment = CENTER
        ws.cell(r, 2, row[0]).alignment = LEFT
        for j in range(3, 9):
            cell = ws.cell(r, j, row[j - 2])
            cell.number_format = "#,##0.00"
            cell.alignment = RIGHT
        for j in range(1, 9):
            ws.cell(r, j).border = BORDER
        r += 1
    # 合计
    ws.cell(r, 2, "合计").font = BOLD
    for j in range(3, 9):
        cell = ws.cell(r, j, round(totals[j - 3], 2))
        cell.number_format = "#,##0.00"
        cell.alignment = RIGHT
        cell.font = BOLD
    for j in range(1, 9):
        ws.cell(r, j).border = BORDER

    # 列宽 + 打印一页
    for j, w in enumerate([4.4, 8.8, 16.3, 16.9, 16.8, 18.8, 19.3, 16.3], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3)


def _uncollected_rows(conn, year: int, month: int):
    """某年度开票的未收明细（发票级，按经办人份额；红冲原票剔除）

    Returns: (rows, total, person_tot)
    rows: [[name, 'X月', buyer, uncol], ...]
    total: 该年未收合计
    person_tot: {name: 该人未收合计}
    """
    # 红冲映射（原票 → {经办人: 份额}）
    red_by_orig: Dict[str, Dict[str, float]] = {}
    reds = conn.execute(
        "SELECT invoice_no, orig_invoice_no FROM invoice "
        "WHERE total_amount<0 AND substr(invoice_date,1,4)=? AND orig_invoice_no IS NOT NULL",
        (str(year),)).fetchall()
    for red in reds:
        for c in conn.execute("SELECT person_name, billing_amount FROM charge_detail WHERE invoice_no=?",
                              (red["invoice_no"],)).fetchall():
            red_by_orig.setdefault(red["orig_invoice_no"], {})[c["person_name"]] = abs(c["billing_amount"])

    invs = conn.execute(
        "SELECT invoice_no, invoice_date, buyer FROM invoice "
        "WHERE total_amount>=0 AND substr(invoice_date,1,4)=?"
        " AND EXISTS (SELECT 1 FROM charge_detail cd WHERE cd.invoice_no=invoice.invoice_no)"
        " ORDER BY invoice_date, invoice_no", (str(year),)).fetchall()
    rows = []
    total = 0.0
    person_tot: Dict[str, float] = {}
    for inv in invs:
        cds = conn.execute(
            "SELECT person_name, billing_amount FROM charge_detail WHERE invoice_no=? ORDER BY id",
            (inv["invoice_no"],)).fetchall()
        if not cds:
            continue
        rem = {c["person_name"]: c["billing_amount"] for c in cds}
        got: Dict[str, float] = {}
        for rec in conn.execute("SELECT amount FROM collection WHERE invoice_no=? ORDER BY id",
                                (inv["invoice_no"],)).fetchall():
            g = allocate_receipt(rem, rec["amount"])
            for name, val in g.items():
                got[name] = got.get(name, 0.0) + val
        mon = f"{int(inv['invoice_date'].split('-')[1])}月"
        for name, billing in rem.items():
            eff = max(billing - red_by_orig.get(inv["invoice_no"], {}).get(name, 0.0), 0.0)
            uncol = round(eff - got.get(name, 0.0), 2)
            if uncol > 0.01:
                rows.append([name, mon, inv["buyer"] or "", uncol])
                total += uncol
                person_tot[name] = round(person_tot.get(name, 0.0) + uncol, 2)
    return rows, round(total, 2), person_tot


def _write_uncollected_sheet(ws, year: int, month: int) -> None:
    conn = get_conn()
    try:
        rows, total, person_tot = _uncollected_rows(conn, year, month)
    finally:
        conn.close()

    ws.merge_cells("A1:E1")
    ws["A1"] = "浙江震天律师事务所"
    ws["A1"].font = Font(size=13, bold=True)
    ws["A1"].alignment = CENTER
    ws.merge_cells("A2:E2")
    ws["A2"] = f"{year}年度开具发票截止{year}年{month}月律师未收款明细表"
    ws["A2"].alignment = CENTER

    for c, h in enumerate(["姓名", "发票开具月份", "对方抬头", "未收款金额", "合计"], 1):
        cell = ws.cell(3, c, h)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = BORDER

    r = 4
    prev_name = None
    for name, mon, buyer, uncol in rows:
        if name != prev_name:
            ws.cell(r, 1, name).alignment = LEFT
        ws.cell(r, 2, mon).alignment = CENTER
        ws.cell(r, 3, buyer).alignment = LEFT
        cell = ws.cell(r, 4, uncol)
        cell.number_format = "#,##0.00"
        cell.alignment = RIGHT
        if name != prev_name:
            cell = ws.cell(r, 5, person_tot.get(name, 0.0))
            cell.number_format = "#,##0.00"
            cell.alignment = RIGHT
        for c in range(1, 6):
            ws.cell(r, c).border = BORDER
        prev_name = name
        r += 1
    # 合计
    ws.cell(r, 1, "合计").font = BOLD
    cell = ws.cell(r, 4, round(total, 2))
    cell.number_format = "#,##0.00"
    cell.alignment = RIGHT
    cell.font = BOLD
    for c in range(1, 6):
        ws.cell(r, c).border = BORDER

    for j, w in enumerate([8, 10, 40, 13, 13], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3)


def _detail_sheets(year: int, month_to: int) -> list:
    """未收款明细 sheet 列表：[('未收款明细2025年01月', 2025, 1), ('未收款明细2024年12月', 2024, 12), ...]（年份倒序）"""
    out = [(f"未收款明细{year}年{month_to:02d}月", year, month_to)]
    conn = get_conn()
    try:
        years = [r[0] for r in conn.execute(
            "SELECT DISTINCT substr(invoice_date,1,4) FROM invoice WHERE total_amount>=0 "
            "AND substr(invoice_date,1,4)<? ORDER BY substr(invoice_date,1,4) DESC",
            (str(year),))]
    finally:
        conn.close()
    for y in years:
        out.append((f"未收款明细{y}年12月", int(y), 12))
    return out


def export_invoice_income(out_path: str | Path, year: int, month_to: int) -> Path:
    """生成开票收入表模板表：1~month_to 主表（从新到旧）+ 未收款明细（当年 + 历史年度）"""
    data = build_settlement(year)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for month in range(month_to, 0, -1):
        ws = wb.create_sheet(title=f"{year}{month:02d}")
        _write_main_sheet(ws, year, month, data)
    for title, y, mo in _detail_sheets(year, month_to):
        ws = wb.create_sheet(title=title)
        _write_uncollected_sheet(ws, y, mo)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def export_invoice_income_month(out_path: str | Path, year: int, month: int) -> Path:
    """生成单月开票收入表：当月主表 + 未收款明细"""
    data = build_settlement(year)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(title=f"{year}{month:02d}")
    _write_main_sheet(ws, year, month, data)
    for title, y, mo in _detail_sheets(year, month):
        ws = wb.create_sheet(title=title)
        _write_uncollected_sheet(ws, y, mo)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
